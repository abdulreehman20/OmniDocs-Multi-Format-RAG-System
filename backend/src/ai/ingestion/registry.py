from pathlib import Path

from docx import Document as DocxDocument
from langchain_community.document_loaders import CSVLoader, PyPDFLoader, TextLoader
from langchain_core.documents import Document as LCDocument

from src.ai.ingestion.contracts import ExtractedDocument, ExtractedUnit
from src.config import Settings
from src.db.models import DocumentSourceType


class IngestionParseError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


def _units_from_lc_docs(
    docs: list[LCDocument],
    *,
    page_key: str | None = None,
) -> list[ExtractedUnit]:
    units: list[ExtractedUnit] = []
    for idx, doc in enumerate(docs):
        text = (doc.page_content or "").strip()
        if not text:
            continue
        locator: dict = {"unit_index": idx}
        meta = doc.metadata or {}
        if page_key and page_key in meta:
            locator["page"] = meta[page_key]
            locator["page_start"] = meta[page_key]
            locator["page_end"] = meta[page_key]
        if "source" in meta:
            locator["source"] = meta["source"]
        units.append(ExtractedUnit(text=text, locator=locator))
    return units


def _load_docx(path: Path) -> list[ExtractedUnit]:
    try:
        document = DocxDocument(str(path))
    except Exception as exc:
        raise IngestionParseError(
            "docx_unreadable",
            "DOCX could not be read (corrupt or password-protected).",
        ) from exc

    units: list[ExtractedUnit] = []
    idx = 0
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        units.append(
            ExtractedUnit(
                text=text,
                locator={"unit_index": idx, "kind": "paragraph"},
            )
        )
        idx += 1
    for t_i, table in enumerate(document.tables):
        rows: list[str] = []
        for row in table.rows:
            cells = [((c.text or "").strip()) for c in row.cells]
            rows.append(" | ".join(cells))
        text = "\n".join(r for r in rows if r.strip())
        if not text:
            continue
        units.append(
            ExtractedUnit(
                text=text,
                locator={"unit_index": idx, "kind": "table", "table_index": t_i},
            )
        )
        idx += 1
    return units


def load_document(
    path: str | Path,
    *,
    source_type: str,
    settings: Settings,
) -> ExtractedDocument:
    path = Path(path)
    warnings: list[str] = []
    parser_version = settings.parser_version

    try:
        if source_type in {DocumentSourceType.MARKDOWN, DocumentSourceType.TEXT}:
            docs = TextLoader(str(path), encoding="utf-8").load()
            units = _units_from_lc_docs(docs)
        elif source_type == DocumentSourceType.PDF:
            docs = PyPDFLoader(str(path)).load()
            units = _units_from_lc_docs(docs, page_key="page")
            if not units:
                raise IngestionParseError(
                    "pdf_no_extractable_text",
                    "PDF has no extractable text. Scanned/image-only PDFs are not supported.",
                )
        elif source_type == DocumentSourceType.CSV:
            docs = CSVLoader(str(path)).load()
            units = _units_from_lc_docs(docs)
            if not units:
                raise IngestionParseError("csv_empty", "CSV produced no rows to index.")
        elif source_type == DocumentSourceType.DOCX:
            units = _load_docx(path)
            if not units:
                raise IngestionParseError("docx_empty", "DOCX produced no extractable text.")
        else:
            raise IngestionParseError(
                "unsupported_source_type",
                f"Unsupported source type: {source_type}",
            )
    except IngestionParseError:
        raise
    except Exception as exc:
        raise IngestionParseError("parse_failed", f"Failed to parse document: {exc}") from exc

    if not units:
        raise IngestionParseError("empty_document", "Document produced no extractable text.")

    return ExtractedDocument(
        units=units,
        parser_version=parser_version,
        warnings=warnings,
        source_metadata={"source_type": source_type, "path": path.name},
    )
